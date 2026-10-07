import shutil
from decimal import Decimal
from pathlib import Path
from unittest.mock import patch
from fastapi.testclient import TestClient
import app

def test_workspace_accounting_and_isolation(tmp_path):
    isolated=tmp_path/'test.sqlite3';shutil.copyfile(app.DB,isolated)
    with patch.object(app,'DB',isolated):
        a=TestClient(app.app);b=TestClient(app.app)
        assert a.get('/api/workspace').status_code==401
        assert a.post('/api/register',json={'email':'a@example.test','password':'long-password','name':'Test business','currency':'USD'}).status_code==200
        initial=a.get('/api/workspace').json();bank,wallet=[x['id'] for x in initial['accounts'] if x['type']=='asset']
        assert not initial['transactions']
        from datetime import date
        def tx(kind,amount,account=bank,destination=None):return a.post('/api/transactions',json={'date':str(date.today()),'kind':kind,'description':'Test entry','amount':amount,'account':account,'destination':destination})
        assert tx('funding','1000.00').status_code==200
        assert tx('income','300.00').status_code==200
        assert tx('expense','100.00').status_code==200
        assert tx('transfer','200.00',destination=wallet).status_code==200
        w=a.get('/api/workspace').json();balances={x['id']:x['balance'] for x in w['accounts']}
        assert balances[bank]==100000 and balances[wallet]==20000
        with app.connect() as c:assert c.execute('SELECT SUM(debit)-SUM(credit) FROM entries').fetchone()[0]==0
        assert tx('transfer','10.00',destination=bank).status_code==400
        assert tx('income','0.001').status_code==422
        assert b.post('/api/register',json={'email':'b@example.test','password':'long-password','name':'Other business'}).status_code==200
        assert not b.get('/api/workspace').json()['transactions']
        assert b.post('/api/transactions',json={'date':str(date.today()),'kind':'income','description':'Bad account','amount':'10','account':bank}).status_code==400
        assert a.post('/api/plan',json={'revenue':1000,'expenses':800,'growth':5,'cost_growth':2,'budget':900}).status_code==200
        assert a.get('/api/workspace').json()['plan']['budget']=='900'
        assert a.post('/api/logout').status_code==200
        assert a.get('/api/workspace').status_code==401
        assert a.post('/api/login',json={'email':'a@example.test','password':'wrong-password'}).status_code==401
        assert a.post('/api/login',json={'email':'a@example.test','password':'long-password'}).status_code==200

def test_opening_balances_validation_and_persistence(tmp_path):
    from datetime import date,timedelta
    isolated=tmp_path/'more.sqlite3';shutil.copyfile(app.DB,isolated)
    with patch.object(app,'DB',isolated):
        c=TestClient(app.app)
        creds={'email':'features-test@example.test','password':'valid-password','name':'Feature verification','currency':'INR'}
        assert c.post('/api/register',json=creds).status_code==200
        assert c.post('/api/register',json=creds).status_code==409
        assert c.post('/api/accounts',json={'name':'Savings','opening':'2500.45'}).status_code==200
        w=c.get('/api/workspace').json();savings=next(a for a in w['accounts'] if a['name']=='Savings')
        assert savings['balance']==250045
        assert w['transactions'][0]['kind']=='funding'
        assert c.post('/api/accounts',json={'name':'   ','opening':0}).status_code==400
        assert c.post('/api/accounts',json={'name':'Invalid','opening':-1}).status_code==422
        bank=next(a for a in w['accounts'] if a['type']=='asset')['id']
        assert c.post('/api/transactions',json={'date':str(date.today()+timedelta(days=1)),'kind':'income','description':'Future','amount':5,'account':bank}).status_code==400
        assert c.post('/api/transactions',json={'date':str(date.today()),'kind':'income','description':'   ','amount':5,'account':bank}).status_code==400
        assert c.post('/api/plan',json={'revenue':1000,'expenses':800,'growth':101,'cost_growth':2,'budget':900}).status_code==422
        assert c.post('/api/plan',json={'revenue':1000,'expenses':800,'growth':5,'cost_growth':2,'budget':900}).status_code==200
        c.post('/api/logout');assert c.post('/api/login',json=creds).status_code==200
        assert c.get('/api/workspace').json()['plan']['budget']=='900'
        assert c.post('/api/accounts',json={'name':'Blocked'},headers={'Origin':'https://outside.test'}).status_code==403
        for route in ('/','/dashboard','/transactions','/accounts','/forecast','/budgets'):assert c.get(route).status_code==200

def test_exports_authenticated_content_and_file_integrity(tmp_path):
    import io,zipfile,json
    from openpyxl import load_workbook
    import fitz
    isolated=tmp_path/'exports.sqlite3';shutil.copyfile(app.DB,isolated)
    with patch.object(app,'DB',isolated):
        c=TestClient(app.app)
        assert c.get('/api/export/pdf').status_code==401
        c.post('/api/register',json={'email':'export-audit@example.test','password':'valid-password','name':'Export audit'})
        c.post('/api/accounts',json={'name':'Savings','opening':'1234.56'})
        for fmt in ('pdf','csv','xlsx','json'):
            r=c.get('/api/export/'+fmt);assert r.status_code==200;assert 'attachment' in r.headers['content-disposition']
            if fmt=='pdf':
                pdf=fitz.open(stream=r.content,filetype='pdf');assert len(pdf)>=5;assert '1,234.56' in ''.join(p.get_text() for p in pdf)
                Path('data/export-verification.pdf').write_bytes(r.content)
                pdf[0].get_pixmap(matrix=fitz.Matrix(1.5,1.5)).save('data/export-verification.png')
            if fmt=='csv':
                z=zipfile.ZipFile(io.BytesIO(r.content));assert 'accounts.csv' in z.namelist();assert '1,234.56' in z.read('accounts.csv').decode('utf-8-sig')
            if fmt=='xlsx':assert set(load_workbook(io.BytesIO(r.content)).sheetnames)=={'Summary','Accounts','Transactions','Forecast','Budget','Read me'}
            if fmt=='json':
                data=r.json();assert 'sessions' not in data and 'password' not in r.text and data['data']['business']['name']=='Export audit'
        assert c.get('/api/export/csv?scope=transactions').headers['content-type'].startswith('text/csv')
        assert c.get('/api/export/pdf?scope=wrong').status_code==422
